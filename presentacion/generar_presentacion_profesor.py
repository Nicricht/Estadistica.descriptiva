from pathlib import Path
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.enum.text import PP_ALIGN,MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

OUT=Path(__file__).with_name('Del_Acero_al_Algoritmo_Estilo_Profesor.pptx')
C=lambda h:RGBColor(*h)
NAV=C((14,45,75)); NAV2=C((20,61,94)); TEAL=C((0,173,159)); TDK=C((0,137,154)); BLUE=C((7,117,150)); ORG=C((245,151,83)); RED=C((232,103,74)); GRN=C((58,164,132)); L=C((246,249,252)); EDGE=C((223,230,237)); TXT=C((36,59,84)); MUT=C((89,108,130)); W=C((255,255,255)); BG=C((13,25,42)); FONT='Noto Sans'
prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)

def shape(s,x,y,w,h,fill,line=None,round=True):
 q=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if round else MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h)); q.fill.solid(); q.fill.fore_color.rgb=fill
 if line: q.line.color.rgb=line; q.line.width=Pt(.7)
 else:q.line.fill.background()
 return q

def text(s,x,y,w,h,t,sz=11,col=TXT,b=False,a=PP_ALIGN.LEFT,v=MSO_ANCHOR.MIDDLE):
 q=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); f=q.text_frame; f.clear(); f.word_wrap=True; f.vertical_anchor=v; f.margin_left=f.margin_right=Inches(.02); f.margin_top=f.margin_bottom=Inches(.02); p=f.paragraphs[0]; p.alignment=a; r=p.add_run(); r.text=t; r.font.name=FONT; r.font.size=Pt(sz); r.font.bold=b; r.font.color.rgb=col; return q

def base(title,n,sub=''):
 s=prs.slides.add_slide(prs.slide_layouts[6]); s.background.fill.solid(); s.background.fill.fore_color.rgb=W; shape(s,0,0,13.333,.055,TDK,None,False); shape(s,.58,.48,.045,.42,TEAL,None,False); text(s,.78,.42,11.35,.62,title,25,NAV,True); text(s,12.2,.42,.55,.4,f'{n:02d}',9.5,MUT,True,PP_ALIGN.RIGHT)
 if sub:text(s,.78,1.04,11.4,.38,sub,10.6,MUT)
 text(s,.68,7.17,7.4,.2,'EduBío 360 · Estadística Descriptiva · Biobío 2021',7.3,C((117,132,149))); return s

def card(s,x,y,w,val,label,col=TEAL,sub=''):
 shape(s,x,y,w,1.34,NAV2); text(s,x+.08,y+.15,w-.16,.5,val,19,col,True,PP_ALIGN.CENTER); text(s,x+.08,y+.72,w-.16,.27,label.upper(),8,C((218,228,237)),True,PP_ALIGN.CENTER)
 if sub:text(s,x+.08,y+1.0,w-.16,.22,sub,7.2,C((188,204,216)),False,PP_ALIGN.CENTER)

def note(s,x,y,w,h,lab,body,col=TEAL):
 shape(s,x,y,w,h,L,EDGE); shape(s,x,y,.035,h,col,None,False); text(s,x+.25,y+.12,w-.45,.24,lab,9.8,TXT,True); text(s,x+.25,y+.38,w-.45,h-.46,body,8.9,MUT)

def badge(s,x,y,w,t,col=TEAL):shape(s,x,y,w,.34,col); text(s,x+.04,y+.02,w-.08,.28,t,8.2,W,True,PP_ALIGN.CENTER)

def hbar(s,x,y,w,lab,val,mx,col,rt=None):
 text(s,x,y-.03,2.55,.3,lab,8.8,NAV,True); shape(s,x+2.62,y,w-3.85,.32,C((224,231,238))); shape(s,x+2.62,y,(w-3.85)*val/mx,.32,col); badge(s,x+w-1.05,y,.95,rt or f'{val:.1f}',NAV); text(s,x+w-1.03,y+.02,.91,.28,rt or f'{val:.1f}',8.3,col,True,PP_ALIGN.CENTER)

def table(s,x,y,widths,rows):
 for i,row in enumerate(rows):
  xx=x; head=i==0; fill=NAV if head else (L if i%2 else W); cols=[W]*len(row) if head else [TXT,TDK,MUT]
  for j,(cw,v) in enumerate(zip(widths,row)):
   shape(s,xx,y+i*.5,cw,.5,fill,EDGE,False); text(s,xx+.1,y+i*.5,cw-.2,.5,str(v),9,cols[j],head or j==1,PP_ALIGN.LEFT if j==0 else PP_ALIGN.CENTER); xx+=cw

s=prs.slides.add_slide(prs.slide_layouts[6]); s.background.fill.solid(); s.background.fill.fore_color.rgb=BG; shape(s,0,0,13.333,.07,TEAL,None,False); text(s,.75,.95,11.8,.55,'DEL ACERO AL ALGORITMO',29,W,True); text(s,.77,1.55,10.6,.52,'Estadística Descriptiva aplicada a la Educación Superior del Biobío',15.5,C((207,219,230))); text(s,.77,2.25,11,1.05,'¿Desde qué base de capital humano parte el Biobío para transformar su histórica vocación industrial hacia una economía de manufactura avanzada e Industria 4.0?',15.2,W,True,v=MSO_ANCHOR.TOP); shape(s,.8,4.25,5.45,1.55,C((28,57,88)),C((52,79,107))); text(s,1.05,4.58,1.2,.35,'ACERO',12,ORG,True); text(s,2.1,4.45,3.7,.65,'motores productivos',19,W,True); shape(s,7.05,4.25,5.45,1.55,C((16,72,82)),C((29,103,112))); text(s,7.3,4.58,1.65,.35,'ALGORITMO',12,TEAL,True); text(s,8.72,4.45,3.25,.65,'capacidades transformadoras',17.5,W,True); text(s,.82,6.37,11.3,.4,'Base oficial: Matrículas de Educación Superior del Biobío · 2021',9.3,C((173,192,207)))

s=base('Contexto del análisis y muestra',2,'Una misma base exige unidades de análisis distintas según la pregunta.')
for x,v,l,su,c in [(0.78,'106.555','base original','28 columnas',TEAL),(3.57,'101.093','matrículas de pregrado','análisis de personas',TEAL),(6.36,'101.067','pregrado con precio','arancel > $0',TEAL),(9.15,'1.273','ofertas únicas','comparación de precios',ORG)]:card(s,x,1.72,2.55,v,l,c,su)
shape(s,.78,3.42,5.8,2.55,L,EDGE); text(s,1.02,3.68,5.35,.34,'¿Qué hacemos con la base?',13.5,NAV,True)
for i,t in enumerate(['Filtramos las matrículas de pregrado.','Separamos los registros con arancel válido.','Para comparar precios, dejamos una sola vez cada oferta académica.']):badge(s,1.02,4.2+i*.52,.42,str(i+1)); text(s,1.58,4.13+i*.52,4.7,.38,t,9.3,TXT)
note(s,6.88,3.42,4.82,2.55,'DECISIÓN METODOLÓGICA','No repetimos el mismo arancel una vez por cada estudiante. Las 1.273 ofertas únicas permiten comparar precios sin que una carrera con más matrícula pese artificialmente más.',ORG)
text(s,.78,6.28,11.92,.34,'Control de calidad: 0 duplicados completos · nulos solo en acreditación · 26 registros con arancel $0 excluidos del análisis de precios.',8.5,MUT,True,PP_ALIGN.CENTER)

s=base('Áreas de estudio en frecuencias relativas',3,'Los porcentajes permiten comparar el peso de cada área dentro de las 101.093 matrículas de pregrado.'); table(s,.72,1.66,[4.7,2,2.2],[('Área del conocimiento','Participación','Lectura'),('Tecnología','27,6%','Mayor participación'),('Salud','24,3%','Segunda mayor'),('Administración y Comercio','13,4%','Participación intermedia'),('Educación','11,4%','Participación intermedia'),('Otros campos','23,3%','Resto de las áreas')]); note(s,.72,5.15,8.9,1.08,'HALLAZGO PRINCIPAL','Tecnología y Salud concentran 51,9% de la matrícula. Las cuatro áreas principales representan 76,7% del total.'); card(s,10.02,1.84,2.4,'51,9%','Tecnología + Salud'); card(s,10.02,3.38,2.4,'76,7%','4 áreas principales',ORG)

s=base('Distribución por género',4,'Mostramos el total y algunas diferencias dentro de familias vinculadas al análisis regional.'); card(s,.78,1.68,2.5,'54,3%','femenino'); card(s,3.55,1.68,2.5,'45,7%','masculino',BLUE); text(s,6.45,1.62,5.5,.34,'Participación femenina dentro de familias seleccionadas',11,NAV,True)
for i,(l,v,c) in enumerate([('Industria y manufactura',17.17,ORG),('Digital y TIC',11.5,TEAL),('Automatización y robótica',5.79,BLUE)]): hbar(s,6.45,2.18+i*.75,6.0,l,v,100,c,f'{v:.1f}%')
note(s,.78,3.55,5.27,1.6,'LECTURA DESCRIPTIVA','En el total hay más mujeres que hombres. Sin embargo, algunas familias tecnológicas presentan una participación femenina mucho menor.'); note(s,6.45,4.75,5.3,1.2,'LÍMITE DE LA BASE','Podemos describir la diferencia. La base no permite afirmar por qué ocurre ni atribuirla a la estructura productiva.',ORG)

s=base('Edad promedio según tipo de institución',5,'Los institutos profesionales y CFT presentan edades medias mayores que las universidades.'); shape(s,.72,1.62,11.9,3.1,L,EDGE)
for i,(l,v,c) in enumerate([('Universidades CRUCH',22.8,BLUE),('Universidades Privadas',24.0,TDK),('Centros de Formación Técnica (CFT)',25.7,GRN),('Institutos Profesionales (IP)',26.5,ORG)]):hbar(s,1.02,1.97+i*.64,10.9,l,v,30,c,f'{v:.1f} a.')
note(s,.72,5.03,11.9,1.1,'ANÁLISIS DEMOGRÁFICO','La edad promedio aumenta desde las universidades CRUCH hacia CFT e IP. Esto describe una diferencia entre tipos de institución; no identifica su causa.')

s=base('Del acero al algoritmo: clasificación propia',6,'Agrupamos carreras en nueve familias para construir una lectura regional del capital humano.'); shape(s,.72,1.62,5.82,4.55,L,EDGE); shape(s,6.78,1.62,5.82,4.55,L,EDGE); text(s,1.02,1.89,2,.42,'ACERO',15,ORG,True); text(s,3.02,1.9,3.1,.36,'Motores productivos',11,NAV,True,PP_ALIGN.RIGHT)
for i,t in enumerate(['Industria y manufactura','Construcción e infraestructura','Logística y puertos','Forestal y madera','Pesca y acuicultura','Agroalimentario']):badge(s,1.03,2.48+i*.53,.43,str(i+1),ORG); text(s,1.62,2.42+i*.53,4.45,.34,t,9.2,TXT,True)
text(s,7.08,1.89,2.4,.42,'ALGORITMO',15,TEAL,True); text(s,9.48,1.9,2.7,.36,'Capacidades transformadoras',10.5,NAV,True,PP_ALIGN.RIGHT)
for i,t in enumerate(['Digital y TIC','Automatización y robótica','Energía, sustentabilidad y biotecnología']):badge(s,7.08,2.62+i*.74,.43,str(i+7)); text(s,7.68,2.54+i*.74,4.15,.45,t,9.7,TXT,True)
note(s,6.98,4.96,5.32,.88,'REGLA','Es una taxonomía creada para este trabajo. Cada carrera se asigna a una sola familia para evitar doble conteo.')

s=base('Hallazgo central: motores productivos vs. capacidades transformadoras',7,'La comparación resume la composición de la matrícula de pregrado en 2021.'); card(s,.9,1.7,3.15,'20.368','motores productivos',ORG,'20,15%'); card(s,4.32,1.7,3.15,'7.000','capacidades transformadoras',TEAL,'6,92%'); card(s,8.55,1.7,3.05,'≈ 2,9 : 1','relación entre bloques',GRN,'productivas / transformadora'); text(s,.95,3.58,10.8,.34,'Participación dentro de las 101.093 matrículas de pregrado',10.7,NAV,True)
for y,v,c,l in [(4.18,20.15,ORG,'20,15%'),(4.92,6.92,TEAL,'6,92%')]: shape(s,.95,y,9.8,.48,C((225,232,238))); shape(s,.95,y,9.8*v/25,.48,c); text(s,11,y-.06,1.25,.55,l,11,c,True,PP_ALIGN.RIGHT)
note(s,.95,5.75,11.3,.78,'INTERPRETACIÓN','En esta fotografía de 2021, el bloque productivo pesa cerca de tres veces más que el transformador. Esto no demuestra déficit de profesionales.')

s=base('Pregunta 1 · ¿Hay áreas donde las carreras sean más caras?',8,'Comparamos la mediana del arancel sobre 1.273 ofertas académicas únicas.'); text(s,.85,1.6,7.65,.34,'Mediana de arancel por área (millones de pesos)',10.7,NAV,True)
for i,(l,v,c) in enumerate([('Derecho',3.586,ORG),('Ciencias Básicas',3.29,GRN),('Agropecuaria',2.982,TEAL),('Tecnología',2.122,BLUE),('Salud',2.08,TDK)]): hbar(s,.88,2.08+i*.64,8.3,l,v,4,c,f'${v:.3f}M')
card(s,9.55,1.8,2.45,'$2.030M','mediana global'); note(s,9.2,3.45,3.15,2.08,'RESPUESTA','Sí. Derecho, Ciencias Básicas y Agropecuaria presentan medianas por encima de la mediana global. Usamos mediana porque los valores extremos la afectan menos.',ORG)

s=base('Pregunta 2 · Edad y tipo de institución',9,'Comparamos la composición institucional de dos grupos de edad con porcentajes.')
for x,t,vals in [(0.78,'Estudiantes de 15 a 19 años',[('Universidad CRUCH',46.24,TEAL),('Instituto Profesional',19.3,BLUE)]),(6.85,'Estudiantes de 40 años o más',[('Instituto Profesional',52.93,ORG),('Universidad CRUCH',10.88,TEAL)])]:
 shape(s,x,1.72,5.55,3.05,L,EDGE); text(s,x+.28,1.98,4.95,.38,t,12.5,NAV,True)
 for i,(l,v,c) in enumerate(vals):hbar(s,x+.28,2.62+i*.83,4.95,l,v,60,c,f'{v:.1f}%')
note(s,.78,5.12,11.62,1.1,'INTERPRETACIÓN','La composición institucional cambia entre grupos de edad. Observamos una asociación descriptiva, pero la base no permite afirmar que la edad cause la elección de institución.')

s=base('Pregunta 3 · ¿Qué aranceles están entre los más altos?',10,'La pauta pide diseñar un criterio reproducible; usamos el percentil 90 (P90), trabajado en el Laboratorio 4.')
for x,w,v,l,c,su in [(0.82,2.65,'$4.106.400','percentil 90',TEAL,'P90'),(3.78,2.35,'128','ofertas sobre P90',RED,'de 1.273'),(6.45,2.35,'10,1%','del total',GRN,''),(9.12,2.8,'$8.783.670','valor máximo',ORG,'')]:card(s,x,1.72,w,v,l,c,su)
text(s,.84,3.64,11,.35,'Cómo leer el P90',11,NAV,True); shape(s,.84,4.23,10.68,.54,C((223,231,238))); shape(s,.84,4.23,9.61,.54,TEAL); shape(s,10.45,4.23,1.07,.54,ORG); text(s,1.2,4.26,8.8,.4,'≈ 90% de las ofertas está en o bajo $4.106.400',9.3,W,True,PP_ALIGN.CENTER); text(s,10.52,4.26,.92,.4,'≈10%',9,W,True,PP_ALIGN.CENTER); note(s,.84,5.24,11.1,1.03,'RESPUESTA','El P90 es $4.106.400. Lo superan 128 ofertas, equivalentes al 10,1% del total. Ese es el grupo que analizamos como arancel alto.')

s=base('¿Qué caracteriza al grupo de ofertas sobre P90?',11,'Describimos las 128 ofertas de arancel alto sin interpretar estas variables como causas del precio.')
for i,(t,a,b,c) in enumerate([('TIPO DE INSTITUCIÓN','73 CRUCH','55 privadas',TEAL),('ÁREA DEL CONOCIMIENTO','38 Salud','38 Tecnología',GRN),('PROVINCIA','123 Concepción','5 Biobío',ORG),('DURACIÓN MEDIANA','10 semestres','resto: 5',BLUE)]):
 x=.78+(i%2)*6.05; y=1.78+(i//2)*1.72; shape(s,x,y,5.55,1.42,L,EDGE); badge(s,x+.22,y+.21,1.85,t,c); text(s,x+.25,y+.67,2.35,.45,a,14,NAV,True); text(s,x+2.7,y+.67,2.55,.45,b,14,c,True,PP_ALIGN.RIGHT)
note(s,.78,5.42,11.62,.88,'LECTURA','Las ofertas altas se concentran en determinados tipos de institución, áreas y territorio, y tienen mayor duración mediana. La base no demuestra que estas variables causen el precio.')

s=base('Capacidades transformadoras por provincia',12,'La proporción y la cantidad absoluta cuentan historias distintas.'); prov=[('ARAUCO','8,55%','232 matrículas',ORG,232),('BIOBÍO','7,61%','966 matrículas',GRN,966),('CONCEPCIÓN','6,77%','5.802 matrículas',TEAL,5802)]
for i,(n,pct,nn,c,cnt) in enumerate(prov):
 x=.82+i*4.05; shape(s,x,1.78,3.52,2.32,L,EDGE); text(s,x+.22,2.03,3.05,.35,n,11,NAV,True,PP_ALIGN.CENTER); text(s,x+.25,2.52,3,.62,pct,24,c,True,PP_ALIGN.CENTER); text(s,x+.22,3.3,3.08,.34,nn,10,MUT,True,PP_ALIGN.CENTER); size=.42+.85*(cnt/5802)**.5; q=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(2.58+i*4.05-size/2),Inches(4.78-size/2),Inches(size),Inches(size)); q.fill.solid(); q.fill.fore_color.rgb=c; q.line.fill.background()
note(s,.82,5.45,11.62,.88,'HALLAZGO TERRITORIAL','Arauco tiene la mayor proporción dentro de su provincia, pero Concepción concentra por lejos el mayor volumen absoluto de matrículas transformadoras.')

s=base('Conclusiones principales',13,'Tres resultados responden el objetivo del trabajo y una cuarta idea delimita lo que no podemos afirmar.')
for i,(num,t,bod,tag,c) in enumerate([('1','COMPOSICIÓN REGIONAL','20,15% productivo vs. 6,92% transformador','≈ 2,9 : 1',ORG),('2','PRECIOS Y EDAD','Hay áreas con arancel mediano mayor y cambia la composición institucional por edad.','',TEAL),('3','ARANCELES ALTOS','128 de 1.273 ofertas superan el P90 de $4.106.400.','10,1%',GRN)]):
 x=.78+i*4.05; shape(s,x,1.74,3.55,3.18,L,EDGE); badge(s,x+.22,1.98,.5,num,c); text(s,x+.82,1.95,2.42,.42,t,10.5,NAV,True); text(s,x+.24,2.66,3.08,1.15,bod,9.8,TXT,True,v=MSO_ANCHOR.TOP)
 if tag:badge(s,x+.88,4.23,1.8,tag,c)
note(s,.78,5.3,11.62,1.05,'LÍMITE DEL ESTUDIO','La base corresponde solo a 2021. Describe diferencias y asociaciones, pero no demuestra causalidad, déficit profesional ni evolución temporal.',RED)

s=base('Cierre · Acero + Algoritmo',14,'La presentación termina donde comienza una nueva pregunta de investigación.'); shape(s,.9,1.72,11.54,3.05,NAV2); text(s,1.25,2.08,10.84,.42,'PREGUNTA FUTURA',10.5,TEAL,True,PP_ALIGN.CENTER); text(s,1.38,2.7,10.58,1.16,'¿Está evolucionando el capital humano del Biobío al mismo ritmo que la transformación productiva que la región proyecta?',20,W,True,PP_ALIGN.CENTER)
for i,(lab,c) in enumerate([('2021\nLínea base',TEAL),('Datos\ndescriptivos',BLUE),('Acero +\nAlgoritmo',ORG),('Investigación\nfutura',GRN)]):
 x=1+i*3; shape(s,x,5.28,2.35,.92,L,EDGE); text(s,x+.12,5.38,2.11,.66,lab,10,c,True,PP_ALIGN.CENTER)
 if i<3:text(s,x+2.48,5.48,.34,.34,'→',14,MUT,True,PP_ALIGN.CENTER)
text(s,.92,6.58,11.4,.34,'Mensaje final: transformar no significa reemplazar la identidad industrial, sino incorporar más capacidades digitales, automatización, datos e innovación.',9.5,MUT,True,PP_ALIGN.CENTER)

s=base('ANEXO · Herramientas del curso utilizadas',15,'Úsalo si el profesor pregunta qué contenidos de los laboratorios aparecen en el trabajo.'); rows=[('Laboratorio','Contenido aplicado','Uso en el proyecto'),('0 · Pandas','read_excel, head, info, filtros','Carga, revisión y preparación'),('1 · Conceptos','población, muestra, tipos de variables','Definición de unidades de análisis'),('2 · Frecuencias','groupby().size(), relativas, cumsum()','Áreas, género, edad y arancel'),('3 · Gráficos','barras y subplots','Distribuciones y comparación'),('4 · Tendencia y percentiles','mean, median, mode, quantile, crosstab','Preguntas 1, 2 y P90'),('5 · Dispersión','rango, std, CV, agg, isin','Dispersión general y comparaciones')]; table(s,.72,1.58,[1.7,4,5.6],rows)

s=base('ANEXO · Preguntas probables del profesor',16,'Respuestas cortas que debes poder explicar con tus propias palabras.'); qa=[('¿Por qué usamos mediana?','Porque los valores extremos afectan menos a la mediana y representa mejor un valor típico.'),('¿Por qué eliminamos ofertas repetidas?','Para que una carrera con muchos estudiantes no repita su precio cientos de veces.'),('¿Qué significa 2,9 : 1?','Aproximadamente 2,9 matrículas productivas por cada matrícula transformadora.'),('¿Edad causa el tipo de institución?','No. Solo observamos una asociación descriptiva entre los grupos.'),('¿Por qué usamos P90?','Porque necesitábamos un umbral reproducible y P90 usa percentiles trabajados en el Laboratorio 4.'),('¿Qué demuestra el notebook?','Patrones, diferencias y asociaciones en 2021; no causalidad ni evolución temporal.')]
for i,(q,a) in enumerate(qa):
 x=.78+(i%2)*6.05; y=1.62+(i//2)*1.7; shape(s,x,y,5.55,1.4,L,EDGE); text(s,x+.22,y+.16,5.08,.34,q,9.5,NAV,True); text(s,x+.22,y+.55,5.08,.68,a,8.7,MUT)

s=base('ANEXO · Conceptos clave para la defensa',17,'No memorices definiciones largas: entiende qué mide cada concepto y dónde lo usamos.'); con=[('Frecuencia relativa','hi = fi / N × 100','Porcentaje que representa una categoría respecto del total.',TEAL),('Mediana','valor central ordenado','Representa el centro y es menos sensible a valores extremos.',ORG),('Percentil 90','quantile(0.90)','Deja aproximadamente 90% de los datos en o bajo ese valor.',GRN),('Desviación estándar','std()','Mide cuánto se alejan los valores respecto de la media.',BLUE)]
for i,(t,f,d,c) in enumerate(con):
 x=.82+(i%2)*6; y=1.76+(i//2)*2.2; shape(s,x,y,5.48,1.86,L,EDGE); badge(s,x+.22,y+.2,2.15,t,c); text(s,x+.25,y+.72,4.98,.36,f,11.3,NAV,True); text(s,x+.25,y+1.14,4.98,.48,d,8.8,MUT)
note(s,.82,6.08,11.48,.74,'REGLA DE ORO','Resultado observado ≠ causa demostrada. En la defensa, usa “se observa”, “se asocia” o “la base muestra”.',RED)
prs.save(OUT); print('Generado',OUT,len(prs.slides))
