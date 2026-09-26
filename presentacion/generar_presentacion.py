from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pathlib import Path
import zipfile, tempfile, shutil

OUT = Path('presentacion/Del_Acero_al_Algoritmo_FINAL_MORPH.pptx')
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

WHITE=RGBColor(246,249,252); MUTED=RGBColor(181,198,216); NAVY=RGBColor(6,14,26)
CARD=RGBColor(12,29,49); CYAN=RGBColor(36,216,255); ORANGE=RGBColor(255,131,31); GOLD=RGBColor(255,187,76); GREEN=RGBColor(88,214,141); GRAY=RGBColor(82,103,125)
LABS=['PANDAS','VARIABLES','FRECUENCIAS','GRÁFICOS','TENDENCIA','DISPERSIÓN','APLICACIÓN','PREGUNTAS']

def nm(sh,n):
    try: sh.name=n
    except: pass
    return sh

def bg(slide,variant=0):
    r=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,prs.slide_width,prs.slide_height); nm(r,'!!BG')
    r.fill.solid(); r.fill.fore_color.rgb=NAVY; r.line.fill.background()
    # moving accents
    for i,(x,y,w,c) in enumerate([(0.25+variant*.03,.9,2.6,CYAN),(10.1-variant*.02,5.85,2.55,ORANGE),(3.8,6.72,5.2,GRAY)]):
        s=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(.025)); nm(s,f'!!BG_LINE_{i}')
        s.fill.solid(); s.fill.fore_color.rgb=c; s.fill.transparency=25 if i<2 else 55; s.line.fill.background()

def title(slide,kicker,main,sub='',n=0):
    t=slide.shapes.add_textbox(Inches(.55),Inches(.35),Inches(5.4),Inches(.25)); nm(t,'!!KICKER')
    p=t.text_frame.paragraphs[0]; r=p.add_run(); r.text=kicker.upper(); r.font.name='Aptos'; r.font.size=Pt(10); r.font.bold=True; r.font.color.rgb=CYAN
    t=slide.shapes.add_textbox(Inches(.55),Inches(.66),Inches(9.3),Inches(.7)); nm(t,'!!TITLE')
    p=t.text_frame.paragraphs[0]; r=p.add_run(); r.text=main; r.font.name='Aptos Display'; r.font.size=Pt(24); r.font.bold=True; r.font.color.rgb=WHITE
    if sub:
        t=slide.shapes.add_textbox(Inches(.57),Inches(1.3),Inches(9.5),Inches(.43)); nm(t,'!!SUBTITLE')
        p=t.text_frame.paragraphs[0]; r=p.add_run(); r.text=sub; r.font.name='Aptos'; r.font.size=Pt(11); r.font.color.rgb=MUTED
    if n:
        t=slide.shapes.add_textbox(Inches(12.25),Inches(.4),Inches(.4),Inches(.2)); p=t.text_frame.paragraphs[0]; p.alignment=PP_ALIGN.RIGHT
        r=p.add_run(); r.text=f'{n:02d}'; r.font.name='Aptos'; r.font.size=Pt(9); r.font.color.rgb=MUTED

def progress(slide,active):
    x0=.6; gap=1.52; y=7.0
    for i,l in enumerate(LABS):
        c=slide.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x0+i*gap),Inches(y),Inches(.12),Inches(.12)); nm(c,f'!!P{i}')
        c.fill.solid(); c.fill.fore_color.rgb=ORANGE if i==active else (CYAN if i<active else GRAY); c.line.fill.background()
        tb=slide.shapes.add_textbox(Inches(x0+i*gap-.22),Inches(y+.15),Inches(.58),Inches(.15)); p=tb.text_frame.paragraphs[0]; p.alignment=PP_ALIGN.CENTER
        r=p.add_run(); r.text=l; r.font.name='Aptos'; r.font.size=Pt(5.2); r.font.color.rgb=WHITE if i==active else MUTED

def flow(slide,x,y,c=CYAN):
    s=slide.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x),Inches(y),Inches(.26),Inches(.26)); nm(s,'!!FLOW')
    s.fill.solid(); s.fill.fore_color.rgb=c; s.line.color.rgb=WHITE

def card(slide,x,y,w,h,k,v,sub='',c=CYAN,name='card'):
    s=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h)); nm(s,name)
    s.fill.solid(); s.fill.fore_color.rgb=CARD; s.line.color.rgb=c
    t=slide.shapes.add_textbox(Inches(x+.14),Inches(y+.12),Inches(w-.28),Inches(.24)); p=t.text_frame.paragraphs[0]
    r=p.add_run(); r.text=k.upper(); r.font.name='Aptos'; r.font.size=Pt(8); r.font.bold=True; r.font.color.rgb=c
    t=slide.shapes.add_textbox(Inches(x+.14),Inches(y+.42),Inches(w-.28),Inches(.48)); p=t.text_frame.paragraphs[0]
    r=p.add_run(); r.text=str(v); r.font.name='Aptos Display'; r.font.size=Pt(19); r.font.bold=True; r.font.color.rgb=WHITE
    if sub:
        t=slide.shapes.add_textbox(Inches(x+.14),Inches(y+h-.3),Inches(w-.28),Inches(.18)); p=t.text_frame.paragraphs[0]
        r=p.add_run(); r.text=sub; r.font.name='Aptos'; r.font.size=Pt(7.2); r.font.color.rgb=MUTED

def badge(slide,n):
    s=slide.shapes.add_shape(MSO_SHAPE.OVAL,Inches(11.85),Inches(.82),Inches(.58),Inches(.58)); nm(s,'!!Q_BADGE')
    s.fill.solid(); s.fill.fore_color.rgb=ORANGE; s.line.color.rgb=WHITE
    t=slide.shapes.add_textbox(Inches(11.85),Inches(.9),Inches(.58),Inches(.3)); p=t.text_frame.paragraphs[0]; p.alignment=PP_ALIGN.CENTER
    r=p.add_run(); r.text=str(n); r.font.name='Aptos Display'; r.font.size=Pt(14); r.font.bold=True; r.font.color.rgb=WHITE

def bars(slide,x,y,w,h,labels,vals,cols,prefix):
    m=max(vals)*1.1; bw=w/(len(vals)*1.6); gap=(w-len(vals)*bw)/(len(vals)-1)
    for i,(lab,v) in enumerate(zip(labels,vals)):
        bx=x+i*(bw+gap); bh=h*v/m; by=y+h-bh
        s=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(bx),Inches(by),Inches(bw),Inches(bh)); nm(s,f'!!{prefix}{i}')
        s.fill.solid(); s.fill.fore_color.rgb=cols[i]; s.line.fill.background()
        t=slide.shapes.add_textbox(Inches(bx-.05),Inches(by-.28),Inches(bw+.1),Inches(.2)); p=t.text_frame.paragraphs[0]; p.alignment=PP_ALIGN.CENTER
        r=p.add_run(); r.text=f'{v:.2f}' if isinstance(v,float) else str(v); r.font.name='Aptos'; r.font.size=Pt(8); r.font.bold=True; r.font.color.rgb=WHITE
        t=slide.shapes.add_textbox(Inches(bx-.12),Inches(y+h+.07),Inches(bw+.24),Inches(.42)); p=t.text_frame.paragraphs[0]; p.alignment=PP_ALIGN.CENTER
        r=p.add_run(); r.text=lab; r.font.name='Aptos'; r.font.size=Pt(6.5); r.font.color.rgb=MUTED

# 1 portada
s=prs.slides.add_slide(prs.slide_layouts[6]); bg(s,0)
t=slide_title=s.shapes.add_textbox(Inches(.8),Inches(1.0),Inches(5.0),Inches(1.2)); nm(t,'!!TITLE'); tf=t.text_frame
p=tf.paragraphs[0]; r=p.add_run(); r.text='DEL ACERO'; r.font.name='Aptos Display'; r.font.size=Pt(31); r.font.bold=True; r.font.color.rgb=WHITE
p=tf.add_paragraph(); r=p.add_run(); r.text='AL ALGORITMO'; r.font.name='Aptos Display'; r.font.size=Pt(39); r.font.bold=True; r.font.color.rgb=ORANGE
card(s,7.6,2.2,2.1,1.3,'Motores','20,15%','20.368 matrículas',ORANGE,'!!MOTOR')
card(s,10.0,2.2,2.1,1.3,'Transformadoras','6,92%','7.000 matrículas',CYAN,'!!TRANS')
t=s.shapes.add_textbox(Inches(.85),Inches(5.35),Inches(10.8),Inches(.7)); p=t.text_frame.paragraphs[0]
r=p.add_run(); r.text='¿Qué peso tenía en 2021 la formación ligada a los motores productivos frente a las capacidades transformadoras?'; r.font.name='Aptos'; r.font.size=Pt(16); r.font.bold=True; r.font.color.rgb=WHITE
flow(s,6.2,6.0,ORANGE)

# 2 recorrido
s=prs.slides.add_slide(prs.slide_layouts[6]); bg(s,1); title(s,'recorrido','Así está construido el notebook','Pandas → variables → frecuencias → gráficos → tendencia → dispersión → aplicación → preguntas.',2); progress(s,0); flow(s,1.0,5.95)
for i,l in enumerate(LABS):
    row=0 if i<4 else 1; col=i if i<4 else i-4
    card(s,.8+col*3.0,2.15+row*1.65,2.55,1.05,f'ETAPA {i}',l,'pregunta → código → resultado → interpretación',ORANGE if i>=6 else CYAN,f'!!STEP{i}')

# 3 pandas
s=prs.slides.add_slide(prs.slide_layouts[6]); bg(s,2); title(s,'lab 0 · pandas','De 106.555 registros a una base lista para analizar','read_excel(), head(), tail(), shape, info(), value_counts(), filtros y nuevas columnas.',3); progress(s,0); flow(s,3.1,5.95)
for i,(k,v,sub) in enumerate([('Registros','106.555','28 columnas'),('Pregrado','101.093','matrículas'),('Precio > $0','101.067','26 excluidos'),('Ofertas','1.273','unidad precio')]): card(s,.9+i*3.05,2.25,2.55,1.5,k,v,sub,ORANGE if i==3 else CYAN,f'!!KPI{i}')
for i,tv in enumerate(['read_excel()','head()/tail()','shape/info()','value_counts()','isin() + filtros']): card(s,1.0+i*2.35,4.45,2.0,.92,'HERRAMIENTA',tv,'',CYAN if i<3 else ORANGE,f'!!TOOL{i}')

# 4 variables
s=prs.slides.add_slide(prs.slide_layouts[6]); bg(s,3); title(s,'lab 1 · conceptos','Población, muestra y clasificación de variables','La clasificación determina qué tabla, gráfico y medida estadística corresponde.',4); progress(s,1); flow(s,4.5,5.9)
vars=[('GENERO','Nominal',CYAN),('RANGO EDAD','Ordinal',GOLD),('EDAD','Cuantitativa',ORANGE),('ARANCEL','Cuantitativa',ORANGE),('ÁREA','Nominal',CYAN),('DURACIÓN','Discreta',GREEN)]
for i,(a,b,c) in enumerate(vars): card(s,.9+(i%3)*4.05,2.2+(i//3)*1.7,3.45,1.2,a,b,'',c,f'!!VAR{i}')

# 5 frecuencias
s=prs.slides.add_slide(prs.slide_layouts[6]); bg(s,4); title(s,'lab 2 · frecuencias','Contar, ordenar y acumular cuando corresponde','Nominales: fi y hi. Cuantitativas ordenadas o agrupadas: fi, Fi, hi y Hi.',5); progress(s,2); flow(s,6.0,5.9)
for i,(a,b,c) in enumerate([('GÉNERO','54,3% F · 45,7% M',CYAN),('ÁREA','27,6% Tecnología',CYAN),('DURACIÓN','10 sem = 34,6%',GREEN),('EDAD','pd.cut()',ORANGE),('ARANCEL','pd.cut()',ORANGE),('fi · Fi · hi · Hi','según variable',GOLD)]): card(s,.85+(i%3)*4.1,2.15+(i//3)*1.75,3.55,1.25,a,b,'',c,f'!!VAR{i}')

# 6 gráficos
s=prs.slides.add_slide(prs.slide_layouts[6]); bg(s,5); title(s,'lab 3 · gráficos','Las tablas se convierten en imágenes interpretables','Circular, barras, histogramas, subplots y dispersión.',6); progress(s,3); flow(s,7.5,5.9)
bars(s,.8,2.3,4.0,2.6,['Tecnología','Salud','Adm.','Educ.'],[27.6,24.3,13.4,11.4],[CYAN,CYAN,ORANGE,GOLD],'G')
card(s,5.35,2.25,2.2,2.0,'GRÁFICO CIRCULAR','54,3 / 45,7','género',CYAN,'!!PIE')
card(s,7.9,2.25,2.2,2.0,'HISTOGRAMA','edad / arancel','intervalos',ORANGE,'!!HIST')
card(s,10.45,2.25,2.2,2.0,'DISPERSIÓN','edad vs arancel','dos variables',GOLD,'!!SCAT')

# 7 tendencia
s=prs.slides.add_slide(prs.slide_layouts[6]); bg(s,6); title(s,'lab 4 · tendencia y percentiles','¿Dónde está el centro y dónde comienza el grupo superior?','Media, mediana, moda, percentiles, describe(), agg() y crosstab().',7); progress(s,4); flow(s,8.8,5.9)
card(s,.9,2.2,2.4,1.4,'MEDIA','$3,22M','arancel',CYAN,'!!MEAN')
card(s,3.55,2.2,2.4,1.4,'MEDIANA','$3,13M','arancel',ORANGE,'!!MEDIAN')
card(s,6.2,2.2,2.4,1.4,'P50','$3,13M','50% bajo/igual',GOLD,'!!P50')
card(s,8.85,2.2,2.4,1.4,'P90','$5,01M','10% superior',GREEN,'!!P90')
for i,(lab,val) in enumerate([('P25','$1,99M'),('P50','$3,13M'),('P75','$4,26M'),('P90','$5,01M')]): card(s,1.15+i*2.85,4.35,2.35,1.0,lab,val,'',CYAN if i<2 else ORANGE,f'!!P{i}')

# 8 dispersión
s=prs.slides.add_slide(prs.slide_layouts[6]); bg(s,7); title(s,'lab 5 · dispersión','No basta saber el centro: importa cuánto se separan los datos','Rango, desviación estándar, coeficiente de variación y comparación entre grupos.',8); progress(s,5); flow(s,10.1,5.9)
card(s,1.0,2.15,3.1,1.4,'ARANCEL · CV','47,1%','mayor dispersión relativa',ORANGE,'!!CVA')
card(s,4.45,2.15,3.1,1.4,'EDAD · CV','≈25,5%','menor dispersión',CYAN,'!!CVE')
card(s,7.9,2.15,3.1,1.4,'RANGO','máximo − mínimo','medida simple',GOLD,'!!RANGO')
for i,a in enumerate(['Tecnología','Salud','Adm. y Comercio','Educación']): card(s,1.0+i*2.85,4.45,2.4,1.0,'COMPARACIÓN',a,'mean · median · std',CYAN if i<2 else ORANGE,f'!!A{i}')

# 9 aplicación
s=prs.slides.add_slide(prs.slide_layouts[6]); bg(s,8); title(s,'aplicación regional','Del acero al algoritmo','Nueve familias estratégicas se agrupan en motores productivos y capacidades transformadoras.',9); progress(s,6); flow(s,11.5,5.9,ORANGE)
for i,a in enumerate(['Industria','Construcción','Logística','Forestal','Pesca','Agro','Digital/TIC','Automatización','Energía+']): card(s,.8+(i%3)*4.15,2.05+(i//3)*1.3,3.55,.94,'FAMILIA',a,'',ORANGE if i<6 else CYAN,f'!!F{i}')

# 10 hallazgo
s=prs.slides.add_slide(prs.slide_layouts[6]); bg(s,9); title(s,'hallazgo central','El componente productivo es casi tres veces el transformador','Describe composición 2021; no prueba déficit laboral.',10); progress(s,6); flow(s,10.0,5.9,ORANGE)
card(s,2.0,2.2,3.4,2.0,'MOTORES PRODUCTIVOS','20.368','20,15%',ORANGE,'!!MOTOR')
card(s,7.95,2.2,3.4,2.0,'CAPACIDADES TRANSFORMADORAS','7.000','6,92%',CYAN,'!!TRANS')
card(s,5.55,4.85,2.2,1.1,'RELACIÓN','≈ 2,9 : 1','',GOLD,'!!RATIO')

# 11 pregunta 1
s=prs.slides.add_slide(prs.slide_layouts[6]); bg(s,10); title(s,'pregunta 1','¿Hay áreas donde las carreras sean más caras?','Criterio: mediana de arancel de 1.273 ofertas únicas.',11); progress(s,7); flow(s,10.8,5.9,ORANGE); badge(s,1)
bars(s,.9,2.25,10.6,3.25,['Derecho','C. Básicas','Agropec.','Tecnología','Salud','Global'],[3.586,3.290,2.982,2.122,2.080,2.030],[ORANGE,GOLD,GREEN,CYAN,CYAN,GRAY],'Q1')

# 12 pregunta 2
s=prs.slides.add_slide(prs.slide_layouts[6]); bg(s,11); title(s,'pregunta 2','¿Qué cambia con la edad al mirar el tipo de institución?','Comparación descriptiva por grupos etarios e institución.',12); progress(s,7); flow(s,9.5,5.9,ORANGE); badge(s,2)
card(s,1.1,2.15,4.4,1.65,'15–19 AÑOS','46,24% CRUCH','19,30% IP',CYAN,'!!YOUNG')
card(s,7.1,2.15,4.4,1.65,'40+ AÑOS','52,93% IP','10,88% CRUCH',ORANGE,'!!OLD')
card(s,3.35,4.55,6.65,1.25,'INTERPRETACIÓN','La distribución cambia con la edad','asociación, no causalidad',GOLD,'!!Q2I')

# 13 pregunta 3
s=prs.slides.add_slide(prs.slide_layouts[6]); bg(s,12); title(s,'pregunta 3','¿Cuándo un arancel es sustantivamente caro?','Criterio: ofertas por sobre el Percentil 90.',13); progress(s,7); flow(s,8.2,5.9,ORANGE); badge(s,3)
card(s,.85,2.05,2.25,1.35,'P90','$4.106.400','umbral',ORANGE,'!!P90Q')
card(s,3.35,2.05,2.25,1.35,'OFERTAS','128','10,1%',CYAN,'!!N128')
card(s,5.85,2.05,2.25,1.35,'INSTITUCIÓN','73 / 55','CRUCH / privadas',GOLD,'!!INST')
card(s,8.35,2.05,2.25,1.35,'TERRITORIO','123 / 5','Concepción / Biobío',GREEN,'!!TERR')
card(s,2.1,4.25,3.6,1.3,'DURACIÓN','10 sem','resto: 5 sem',ORANGE,'!!DUR')
card(s,7.0,4.25,3.6,1.3,'ÁREAS','Salud 38 · Tecnología 38','',CYAN,'!!AREA')

# 14 cierre
s=prs.slides.add_slide(prs.slide_layouts[6]); bg(s,13); title(s,'cierre','La estadística describe el punto de partida, no inventa causas','El valor del trabajo está tanto en los hallazgos como en reconocer sus límites.',14); progress(s,7); flow(s,6.4,5.9,ORANGE)
card(s,.8,2.2,2.8,1.5,'PRODUCTIVO','20,15%','20.368 matrículas',ORANGE,'!!MOTOR')
card(s,9.75,2.2,2.8,1.5,'TRANSFORMADOR','6,92%','7.000 matrículas',CYAN,'!!TRANS')
card(s,4.65,2.6,4.0,1.65,'CONCLUSIÓN','≈ 2,9 : 1','composición 2021',GOLD,'!!FINAL')
for i,tv in enumerate(['No prueba déficit','No prueba causalidad','No mide salarios/vacantes','No muestra evolución temporal']): card(s,1.0+i*3.0,5.1,2.55,.9,'LÍMITE',tv,'',GRAY,f'!!LIM{i}')

OUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(OUT)

# Agregar Morph por objeto en diapositivas 2..14
tmp=Path(tempfile.mkdtemp(prefix='morph_'))
with zipfile.ZipFile(OUT,'r') as z: z.extractall(tmp)
block='<mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" xmlns:p159="http://schemas.microsoft.com/office/powerpoint/2015/09/main"><mc:Choice Requires="p159"><p:transition spd="slow"><p159:morph option="byObject"/></p:transition></mc:Choice><mc:Fallback><p:transition spd="slow"><p:fade/></p:transition></mc:Fallback></mc:AlternateContent>'
for sf in sorted((tmp/'ppt'/'slides').glob('slide*.xml'), key=lambda p:int(p.stem.replace('slide','')))[1:]:
    xml=sf.read_text(encoding='utf-8')
    pos=xml.rfind('</p:sld>')
    sf.write_text(xml[:pos]+block+xml[pos:],encoding='utf-8')
rep=OUT.with_suffix('.tmp.pptx')
with zipfile.ZipFile(rep,'w',zipfile.ZIP_DEFLATED) as z:
    for f in tmp.rglob('*'):
        if f.is_file(): z.write(f,f.relative_to(tmp))
shutil.move(rep,OUT); shutil.rmtree(tmp,ignore_errors=True)
print(OUT)
