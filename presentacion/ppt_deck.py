from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from ppt_assets import A, OUT, assets

N=RGBColor(5,13,25); C=RGBColor(38,218,255); O=RGBColor(255,132,32); W=RGBColor(246,249,252)
M=RGBColor(180,198,216); G=RGBColor(255,190,80); R=RGBColor(255,91,91); GR=RGBColor(92,112,132); GN=RGBColor(88,214,141)

assets()
prs=Presentation()
prs.slide_width=Inches(13.333)
prs.slide_height=Inches(7.5)
SW,SH=prs.slide_width,prs.slide_height

ST=['BASE','FRECUENCIAS','GRÁFICOS','MEDIDAS','PREGUNTAS','CIERRE']

def nm(x,n):
    try: x.name=n
    except: pass
    return x

def bg(s,k):
    nm(s.shapes.add_picture(str(A/f'bg{k%6}.jpg'),0,0,width=SW,height=SH),'!!BG')
    z=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,SW,SH)
    z.fill.solid(); z.fill.fore_color.rgb=N; z.fill.transparency=23; z.line.fill.background()

def text(s,x,y,w,h,t,size=11,col=W,bold=False,align=PP_ALIGN.LEFT,name=None):
    b=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    nm(b,name or 'text')
    p=b.text_frame.paragraphs[0]
    p.alignment=align
    r=p.add_run(); r.text=t; r.font.name='Aptos'; r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=col
    return b

def title(s,n,k,t,sub=''):
    text(s,.55,.32,5.5,.25,k.upper(),9,C,True,name='!!KICKER')
    text(s,.55,.62,11.4,.72,t,24,W,True,name='!!TITLE')
    if sub: text(s,.57,1.28,11.1,.42,sub,10,M,name='!!SUB')
    text(s,12.2,.4,.45,.2,f'{n:02d}',8,M,True,PP_ALIGN.RIGHT)

def stage(s,a):
    for i,l in enumerate(ST):
        x=.6+i*2.02
        q=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x),Inches(7.02),Inches(.11),Inches(.11))
        q.fill.solid(); q.fill.fore_color.rgb=O if i==a else (C if i<a else GR); q.line.fill.background()
        text(s,x-.32,7.17,.76,.13,l,5,W if i==a else M,False,PP_ALIGN.CENTER)

def flow(s,n):
    q=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(.75+(n*.47)%11.2),Inches(6.08),Inches(.22),Inches(.22))
    q.fill.solid(); q.fill.fore_color.rgb=O if n%3==0 else C; q.line.color.rgb=W

def card(s,x,y,w,h,k,v='',sub='',c=C,name='card',vs=19):
    q=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    nm(q,name); q.fill.solid(); q.fill.fore_color.rgb=RGBColor(10,27,46); q.fill.transparency=6; q.line.color.rgb=c; q.line.transparency=25
    text(s,x+.14,y+.1,w-.28,.2,k.upper(),7,c,True)
    if v: text(s,x+.14,y+.38,w-.28,.48,str(v),vs,W,True)
    if sub: text(s,x+.14,y+h-.29,w-.28,.18,sub,6.6,M)

def pill(s,x,y,w,t,c=C,name='pill'):
    q=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(.36))
    nm(q,name); q.fill.solid(); q.fill.fore_color.rgb=c; q.fill.transparency=76; q.line.color.rgb=c
    text(s,x+.04,y+.075,w-.08,.2,t,7.2,W,True,PP_ALIGN.CENTER)

def note(s,x,y,w,t,c=O):
    q=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(.55))
    nm(q,'!!NOTE'); q.fill.solid(); q.fill.fore_color.rgb=c; q.fill.transparency=80; q.line.color.rgb=c
    text(s,x+.1,y+.12,w-.2,.3,t,8.5,W,True,PP_ALIGN.CENTER)

def footer(s,t='Base de matrículas de Educación Superior del Biobío 2021'):
    text(s,6.6,6.82,6.1,.15,t,5.6,RGBColor(135,152,170),False,PP_ALIGN.RIGHT)

def basic(n,k,ki,ti,su,st):
    s=prs.slides.add_slide(prs.slide_layouts[6]); bg(s,k); title(s,n,ki,ti,su); stage(s,st); flow(s,n); return s

def bars(s,x,y,w,h,labels,vals,cols,maxv=None,prefix='B'):
    m=maxv or max(vals)*1.1
    bw=w/(len(vals)*1.55)
    gap=(w-len(vals)*bw)/(len(vals)-1)
    for i,(lab,v) in enumerate(zip(labels,vals)):
        xx=x+i*(bw+gap); hh=h*v/m
        q=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(xx),Inches(y+h-hh),Inches(bw),Inches(hh))
        nm(q,f'!!{prefix}{i}'); q.fill.solid(); q.fill.fore_color.rgb=cols[i]; q.line.fill.background()
        text(s,xx-.06,y+h-hh-.25,bw+.12,.2,str(v).replace('.',','),7.5,W,True,PP_ALIGN.CENTER)
        text(s,xx-.12,y+h+.05,bw+.24,.35,lab,6,M,False,PP_ALIGN.CENTER)
