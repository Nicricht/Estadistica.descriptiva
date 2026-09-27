from pathlib import Path
from pptx import Presentation
from pptx.util import Inches

PPT = Path('presentacion/Del_Acero_al_Algoritmo_Estilo_Profesor.pptx')
prs = Presentation(PPT)

# La presentación completa tenía 18 diapositivas. Para la exposición dejamos
# solo las que aportan directamente a la historia y la pauta.
# Se eliminan, usando numeración humana original:
# 4 = género (resultado secundario)
# 5 = edad promedio separada (se resume en Pregunta 2)
# 14 = cierre redundante con conclusiones
# 17 = conceptos clave (se cubren en preguntas probables)
# 18 = clasificación de variables (queda en el notebook, como pidió el equipo)
for human_index in sorted([18, 17, 14, 5, 4], reverse=True):
    idx = human_index - 1
    slide_id = prs.slides._sldIdLst[idx]
    prs.part.drop_rel(slide_id.rId)
    del prs.slides._sldIdLst[idx]

assert len(prs.slides) == 13, len(prs.slides)

# Renumerar las diapositivas visibles del estilo del profesor.
# La portada no lleva número. Las demás tienen un pequeño número arriba a la derecha.
for new_number, slide in enumerate(list(prs.slides)[1:], start=2):
    for shape in slide.shapes:
        if not getattr(shape, 'has_text_frame', False):
            continue
        # Posición usada por la función base() del generador.
        if shape.left >= Inches(12.0) and shape.top <= Inches(0.9):
            current = shape.text.strip()
            if current.isdigit() and len(current) <= 2:
                shape.text = f'{new_number:02d}'
                break

prs.save(PPT)
print('Presentación compactada:', len(prs.slides), 'diapositivas')
